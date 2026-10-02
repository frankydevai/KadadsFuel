from openpyxl import Workbook
from zipfile import ZipFile, ZIP_DEFLATED
import re

from dieselup.ingestion.pilot_parser import REQUIRED_COLUMNS, parse_pilot_xls


def test_modern_xlsx_price_file_is_parsed(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.cell(row=5, column=1, value="Account: 12345 - Fleet")
    for column, name in enumerate(REQUIRED_COLUMNS, start=1):
        sheet.cell(row=6, column=column, value=name)
    columns = {name: index + 1 for index, name in enumerate(REQUIRED_COLUMNS)}
    for offset in range(100):
        row = 7 + offset
        values = {
            "Site": 1000 + offset,
            "City": f"City {offset}",
            "ST": "NJ",
            "Prod": "DSL",
            "Rack ID": 1,
            "Rack City": "Rack",
            "Rack ST": "NJ",
            "Cost": 3.0,
            "Total Cost": 3.1,
            "Retail Price": 4.0,
            "Your Price": 3.5,
            "Savings Total": 0.5,
        }
        for name, value in values.items():
            sheet.cell(row=row, column=columns[name], value=value)
    path = tmp_path / "pilot.xlsx"
    workbook.save(path)

    result = parse_pilot_xls(str(path))

    assert result["account_number"] == "12345"
    assert result["row_count"] == 100

    # QuickManage/Excel exports may omit dimensions in an otherwise valid XLSX.
    with ZipFile(path) as source:
        parts={name:source.read(name) for name in source.namelist()}
    parts['xl/worksheets/sheet1.xml']=re.sub(rb'<dimension[^>]*/>',b'',parts['xl/worksheets/sheet1.xml'])
    undimensioned=tmp_path/'no-dimensions.xlsx'
    with ZipFile(undimensioned,'w',ZIP_DEFLATED) as target:
        for name,value in parts.items():target.writestr(name,value)
    recovered=parse_pilot_xls(str(undimensioned))
    assert recovered['account_number']=='12345'
    assert recovered['stops']==result['stops']

    # Header positions can shift when an exporter adds a carrier/logo row.
    sheet.insert_rows(1,3)
    shifted=tmp_path/'shifted-headers.xlsx';workbook.save(shifted)
    assert parse_pilot_xls(str(shifted))['stops']==result['stops']
